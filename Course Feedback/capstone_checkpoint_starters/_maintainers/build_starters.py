"""Build the checkpoint starter folders from one source. (For maintainers.)

    python _maintainers/build_starters.py           # rebuild every checkpoint folder
    python _maintainers/build_starters.py --test    # rebuild, then run each folder's offline tests

Every file lives once, in _maintainers/source/. Each checkpoint folder is generated
from it: the common setup files, plus every library module, sample, and test that was
introduced in that checkpoint or an earlier one. A fix made once in source/ reaches
every folder on the next build, so the folders cannot drift apart.

Do not edit the checkpoint folders by hand: the build replaces them. Run the build with
the Python of an environment that has requirements.txt installed (it needs openpyxl to
write the Excel question sets, and pytest for --test).
"""

from __future__ import annotations

import argparse
import compileall
import csv
import os
import shutil
import subprocess
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
SOURCE = HERE / "source"
STARTERS = HERE.parent
CHECKPOINTS = ["1", "2", "3", "4", "5", "log_analysis"]

# Where each library module first appears. Every file under source/ragkit must be listed.
LIBRARY = {
    "1": ["__init__.py", "config.py", "llm.py", "chat.py", "runlog.py", "pricing.py", "pricing.json",
          "pipelines/__init__.py", "pipelines/base.py", "pipelines/llm_only.py",
          "reports/__init__.py", "reports/common.py", "reports/runs.py"],
    "2": ["corpus.py", "embeddings.py", "retrievers/__init__.py", "retrievers/base.py", "retrievers/bm25.py",
          "retrievers/vector.py", "retrievers/hybrid.py", "retrievers/factory.py", "pipelines/rag.py"],
    "3": ["eval/__init__.py", "eval/questions.py", "eval/recall.py", "eval/judge.py", "eval/run.py",
          "pipelines/oracle.py", "baseline.py", "reports/quality.py"],
    "4": ["retrievers/multistep.py", "retrievers/graph.py"],
    "5": ["structured.py", "retrievers/agentic.py", "pipelines/agent.py"],
    "log_analysis": ["reports/performance.py"],
}
# Samples and tests are assigned by name.
SAMPLES = {"1": "cp1_", "2": "cp2_", "3": "cp3_", "4": "cp4_", "5": "cp5_", "log_analysis": "analyze_logs"}
TESTS = {
    "1": ["__init__.py", "conftest.py", "fakes.py", "test_cp1_"],
    "2": ["fakes_retrieval.py", "test_cp2_", "fixtures/"],
    "3": ["test_cp3_"],
    "4": ["test_cp4_"],
    "5": ["fakes_agents.py", "test_cp5_"],
    "log_analysis": ["test_log_analysis"],
}
DATA_CHECKPOINT = "3"  # the question sets arrive with evaluation

HOW_TO_FILL = [
    ("ID", "A short unique id: GT01, GT02, ... for gt questions; VAL01, ... for validation questions."),
    ("Category", "The kind of question: direct fact, quotation, table lookup, multi-article, explanation, ... "
                 "The quality report shows results per category."),
    ("Question", "Exactly as a user would ask it."),
    ("Reference answer", "The correct answer, copied from the corpus, not from memory."),
    ("Grading notes", "What a PASS must contain, what makes it PARTIAL, what makes it FAIL. The judge reads these."),
    ("Source articles", "The file name(s) that hold the answer, separated by ';', for example "
                        "68th_Academy_Awards.html; 74th_Academy_Awards.html. Retrieval recall is measured against these."),
    ("Split", "gt: used while improving the system. validation: held back, run only to confirm that an "
              "improvement also holds on questions it was not tuned on."),
]


def checkpoint_of(relative: str) -> str | None:
    """The checkpoint that introduces a source file, or None if it is not assigned."""
    parts = relative.split("/", 1)
    if parts[0] == "ragkit":
        return next((cp for cp, files in LIBRARY.items() if parts[1] in files), None)
    if parts[0] == "samples":
        return next((cp for cp, prefix in SAMPLES.items() if parts[1].startswith(prefix)), None)
    if parts[0] == "tests":
        return next((cp for cp, prefixes in TESTS.items() if any(parts[1].startswith(p) for p in prefixes)), None)
    if parts[0] == "data":
        return DATA_CHECKPOINT
    return None


def source_files() -> list[str]:
    files = []
    for folder in ("ragkit", "samples", "tests", "data"):
        for path in sorted((SOURCE / folder).rglob("*")):
            if path.is_file() and "__pycache__" not in path.parts and path.name != ".DS_Store":
                files.append(path.relative_to(SOURCE).as_posix())
    return files


def build(checkpoint: str) -> Path:
    upto = CHECKPOINTS[: CHECKPOINTS.index(checkpoint) + 1]
    target = STARTERS / f"checkpoint_{checkpoint}"
    if target.exists():
        shutil.rmtree(target)
    shutil.copytree(SOURCE / "common", target)  # setup files, identical in every checkpoint

    for relative in source_files():
        if checkpoint_of(relative) not in upto:
            continue
        if relative.startswith("data/") and relative.endswith(".csv"):
            write_workbook(SOURCE / relative, target / relative.replace(".csv", ".xlsx"))
            continue
        destination = target / relative
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(SOURCE / relative, destination)

    shutil.copy2(SOURCE / "guides" / f"checkpoint_{checkpoint}.md", target / "README.md")
    for earlier in upto[:-1]:
        guide = target / "guides" / f"checkpoint_{earlier}.md"
        guide.parent.mkdir(exist_ok=True)
        shutil.copy2(SOURCE / "guides" / f"checkpoint_{earlier}.md", guide)
    return target


def write_workbook(csv_path: Path, xlsx_path: Path) -> None:
    """A question set as an Excel workbook: the questions, and a sheet explaining each column."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    with csv_path.open(encoding="utf-8") as handle:
        rows = list(csv.reader(handle))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Questions"
    for row in rows:
        sheet.append(row)
    widths = {"A": 8, "B": 22, "C": 50, "D": 50, "E": 60, "F": 40, "G": 11}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for cells in sheet.iter_rows():
        for item in cells:
            item.alignment = Alignment(wrap_text=True, vertical="top")
    for item in sheet[1]:
        item.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    guide = workbook.create_sheet("How to fill")
    guide.append(["Column", "What to put in it"])
    for row in HOW_TO_FILL:
        guide.append(list(row))
    guide.column_dimensions["A"].width = 18
    guide.column_dimensions["B"].width = 110
    for item in guide[1]:
        item.font = Font(bold=True)
    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(xlsx_path)


def run_tests(folder: Path) -> bool:
    """Compile every file, then run the folder's own tests, with nothing but the folder on the path."""
    if not compileall.compile_dir(str(folder), quiet=1, force=True):
        return False
    env = {**os.environ, "PYTHONPATH": str(folder)}
    result = subprocess.run([sys.executable, "-m", "pytest", "-p", "no:cacheprovider"],
                            cwd=folder, env=env, capture_output=True, text=True)
    summary = result.stdout.strip().splitlines()[-1] if result.stdout.strip() else result.stderr[-500:]
    print(f"    {summary}")
    for cache in folder.rglob("__pycache__"):
        shutil.rmtree(cache, ignore_errors=True)
    return result.returncode == 0


def main() -> None:
    parser = argparse.ArgumentParser(description="Build the checkpoint starter folders.")
    parser.add_argument("--test", action="store_true", help="run each folder's tests after building it")
    args = parser.parse_args()

    unassigned = [f for f in source_files() if checkpoint_of(f) is None]
    if unassigned:
        raise SystemExit(f"Assign these files to a checkpoint in build_starters.py first: {', '.join(unassigned)}")
    failed = []
    for checkpoint in CHECKPOINTS:
        folder = build(checkpoint)
        count = sum(1 for p in folder.rglob("*") if p.is_file())
        print(f"Built {folder.name}: {count} files")
        if args.test and not run_tests(folder):
            failed.append(folder.name)
    if failed:
        raise SystemExit(f"Tests failed in: {', '.join(failed)}")


if __name__ == "__main__":
    main()
